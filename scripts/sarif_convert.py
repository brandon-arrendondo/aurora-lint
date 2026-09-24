#!/usr/bin/env python3
"""Convert an aurora-lint SARIF report to CSV or XLSX.

aurora-lint writes one full-fidelity report format, SARIF 2.1.0
(`aurora-lint SRC --export report.sarif`). Spreadsheets are a view of that
report, derived here rather than built into the tool; the report carries
everything a row needs (rule description, source line, file hash), so this
runs without the scanned tree or the binary.

    python scripts/sarif_convert.py report.sarif findings.csv
    python scripts/sarif_convert.py report.sarif findings.xlsx   # needs openpyxl

One row per active finding, in the tracker-import layout the tool's former
CSV/XLSX export used, plus a Tags column naming the tool that produced the
report. Suppressed findings are left out unless --include-suppressed.
"""

import argparse
import csv
import json
import sys
from pathlib import Path

HEADERS = [
    "Title",
    "Description",
    "Work Item Type",
    "State",
    "Severity",
    "Priority",
    "Tags",
]
# Fixed values the former export wrote on every row; kept so an import
# mapping built against it still lines up.
ROW_CONSTANTS = ["Bug", "Proposed", "1 - Critical", "1"]
REVIEW_NOTE = (
    " (aurora-lint could not confidently determine this is a violation"
    " -- review before acting.)"
)


def sarif_rows(sarif: dict, include_suppressed: bool = False) -> list[list[str]]:
    """Build one spreadsheet row per result across every run in `sarif`."""
    rows = []
    for run in sarif.get("runs", []):
        driver = run.get("tool", {}).get("driver", {})
        tool_name = driver.get("name", "")
        descriptions = {
            rule.get("id"): rule.get("shortDescription", {}).get("text", "Unknown rule")
            for rule in driver.get("rules", [])
        }
        artifacts = run.get("artifacts", [])

        for result in run.get("results", []):
            if result.get("suppressions") and not include_suppressed:
                continue
            rule_id = result.get("ruleId", "")
            location = result.get("locations", [{}])[0].get("physicalLocation", {})
            artifact_location = location.get("artifactLocation", {})
            region = location.get("region", {})

            file_hash = "unknown"
            index = artifact_location.get("index")
            if isinstance(index, int) and 0 <= index < len(artifacts):
                sha = artifacts[index].get("hashes", {}).get("sha-256")
                if sha:
                    file_hash = sha[:8]

            needs_review = result.get("properties", {}).get("requiresManualReview", False)
            title = "{}{}:{}:{} version:{}".format(
                "[NEEDS MANUAL REVIEW] " if needs_review else "",
                rule_id,
                artifact_location.get("uri", ""),
                region.get("startLine", ""),
                file_hash,
            )
            description = "{} - {}: {}{}".format(
                rule_id,
                descriptions.get(rule_id, "Unknown rule"),
                region.get("snippet", {}).get("text", "(line not found)"),
                REVIEW_NOTE if needs_review else "",
            )
            rows.append([title, description, *ROW_CONSTANTS, tool_name])
    return rows


def write_csv(rows: list[list[str]], path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(HEADERS)
        writer.writerows(rows)


def write_xlsx(rows: list[list[str]], path: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        sys.exit("XLSX output needs openpyxl: pip install openpyxl (or write .csv)")

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9D9D9")
    for row in rows:
        sheet.append(row)
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = min(width + 2, 100)
    workbook.save(path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Convert an aurora-lint SARIF report to CSV or XLSX."
    )
    parser.add_argument("sarif", type=Path, help="SARIF report from --export")
    parser.add_argument("output", type=Path, help="output path ending in .csv or .xlsx")
    parser.add_argument(
        "--include-suppressed",
        action="store_true",
        help="also emit findings carrying an in-source suppression",
    )
    args = parser.parse_args(argv)

    suffix = args.output.suffix.lower()
    if suffix not in (".csv", ".xlsx"):
        parser.error(f"output must end in .csv or .xlsx, not '{suffix or args.output}'")

    sarif = json.loads(args.sarif.read_text(encoding="utf-8"))
    rows = sarif_rows(sarif, include_suppressed=args.include_suppressed)
    (write_csv if suffix == ".csv" else write_xlsx)(rows, args.output)
    print(f"Wrote {len(rows)} findings to {args.output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
